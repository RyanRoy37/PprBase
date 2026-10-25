"""Turns a format + :class:`PaperState` into a fully formatted ``.docx``.

The engine only reads from state. For an imported Word template it loads a
fresh copy, keeps its header and page setup, and replaces the question body.

Layout rules implemented here:

* Header block: school name, "Subject — Class" line, duration / total marks line.
* Instructions block under the header.
* One section heading per section (bold, centred, with a top border rule).
* Question numbering ``1.``, ``2.`` ... and sub-parts ``(i)``, ``(a)`` etc.
* MCQ options laid out ``(a) ... (b) ...`` on indented lines.
* Marks right-aligned on the question line using a right-aligned tab stop.
"""

from __future__ import annotations

import os
import re
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph

from core import image_handler, table_builder
from core.format_loader import section_instruction
from core.state_manager import PaperState

OPTION_LABELS = ["a", "b", "c", "d", "e", "f", "g", "h"]


class DocumentEngine:
    """Builds a Word document for a given format and paper state."""

    def __init__(self, fmt: dict[str, Any], state: PaperState) -> None:
        self.fmt = fmt
        self.state = state
        self._doc: Document | None = None
        self._right_tab_cm = 16.0  # right margin for the marks tab stop
        self._template_heading_pt = 13.0
        self._pending_template_duration = ""
        self._temp_images: list[image_handler.PreparedImage] = []

    # -- public api -------------------------------------------------------
    def build_document(self) -> Document:
        """Generate and return a complete :class:`docx.Document`."""
        for prepared in self._temp_images:
            prepared.cleanup()
        self._temp_images.clear()
        template = self.fmt.get("docx_template")
        if template:
            if not Path(template).is_file():
                raise FileNotFoundError(f"The Word template is missing: {template}")
            self._doc = Document(template)
            self._prepare_template()
            self._add_instructions()
        else:
            self._doc = Document()
            self._configure_base_styles()
            self._add_header()
            self._add_instructions()
        for section in self.fmt.get("sections", []):
            self._add_section(section)
        return self._doc

    def _prepare_template(self) -> None:
        """Clear sample questions while retaining the original page and header."""
        body = self._doc.element.body
        start = self.fmt.get("template_body_start")
        if start is None or not 0 <= start < len(body):
            raise ValueError("This Word template has no valid question start. Import it again.")
        self._update_template_header(start)
        font_sizes: Counter[float] = Counter()
        first_body_paragraph = None
        for element in list(body)[start:]:
            if element.tag != qn("w:p"):
                continue
            paragraph = Paragraph(element, self._doc)
            if not paragraph.text.strip():
                continue
            if first_body_paragraph is None:
                first_body_paragraph = paragraph
            for run in paragraph.runs:
                if run.font.size and run.text.strip():
                    font_sizes[round(run.font.size.pt, 1)] += len(run.text.strip())
        if font_sizes:
            self._doc.styles["Normal"].font.size = Pt(font_sizes.most_common(1)[0][0])
        if first_body_paragraph:
            first_size = next((run.font.size.pt for run in first_body_paragraph.runs if run.font.size), None)
            if first_size:
                self._template_heading_pt = first_size
        for element in list(body)[start:]:
            if element.tag != qn("w:sectPr"):
                body.remove(element)
        # Old question illustrations should not travel inside a new paper.
        used = {
            value for node in body.iter() for key, value in node.attrib.items()
            if key in (qn("r:embed"), qn("r:id"))
        }
        for rel_id, rel in list(self._doc.part.rels.items()):
            if rel.reltype == RT.IMAGE and rel_id not in used:
                self._doc.part.drop_rel(rel_id)
        section = self._doc.sections[0]
        self._right_tab_cm = max(1.0, section.page_width.cm - section.left_margin.cm - section.right_margin.cm - 0.2)
        if self._pending_template_duration:
            self._doc.add_paragraph(f"Time: {self._pending_template_duration}")

    def _update_template_header(self, body_start: int) -> None:
        """Apply edits to detected header fields without rebuilding the header layout."""
        paragraphs = [
            Paragraph(element, self._doc)
            for element in list(self._doc.element.body)[:body_start]
            if element.tag == qn("w:p")
        ]
        first = next((paragraph for paragraph in paragraphs if paragraph.text.strip()), None)
        school = self.fmt.get("school_name", "").strip()
        if first and school and first.text.strip() != school:
            self._replace_paragraph_text(first, school)
        title = self.fmt.get("paper_title", "").strip()
        duration = self.fmt.get("duration", "").strip()
        found_duration = False
        for paragraph in paragraphs:
            original = paragraph.text
            if paragraph is not first and title and re.search(r"\b(?:test|exam|assessment|paper)\b", original, re.I):
                if original.strip() != title:
                    self._replace_paragraph_text(paragraph, title)
                break
        for paragraph in paragraphs:
            original = paragraph.text
            updated = original
            for field, pattern in (
                ("subject", r"(\bsubject\s*[:\-]\s*)(.*?)(?=\s{2,}|\t|\bclass\s*[:\-]|\bmarks?\s*[:\-]|$)"),
                ("class", r"(\bclass\s*[:\-]\s*)([a-z0-9]+)"),
            ):
                value = str(self.fmt.get(field, "")).strip()
                match = re.search(pattern, updated, re.I)
                if value and match and match.group(2).strip() != value:
                    updated = updated[:match.start(2)] + value + updated[match.end(2):]
            marks = self._fmt_marks(self.fmt.get("total_marks"))
            match = re.search(r"(\bmarks?\s*[:\-]\s*)(\d+(?:\.\d+)?)", updated, re.I)
            if match and match.group(2) != marks:
                updated = updated[:match.start(2)] + marks + updated[match.end(2):]
            match = re.search(r"(\b(?:time|duration)\s*[:\-]\s*)([^\t]+)", updated, re.I)
            if match:
                found_duration = True
                if duration and match.group(2).strip() != duration:
                    updated = updated[:match.start(2)] + duration + updated[match.end(2):]
            if updated != original:
                self._replace_paragraph_text(paragraph, updated)
        self._pending_template_duration = duration if duration and not found_duration else ""

    @staticmethod
    def _replace_paragraph_text(paragraph: Paragraph, value: str) -> None:
        """Keep the paragraph's alignment and first run styling when text changes."""
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(value)

    def save(self, output_path: str) -> str:
        """Build (if needed) and write the document to ``output_path``.

        Returns the path written. Cleans up any temporary converted images.
        """
        try:
            if self._doc is None:
                self.build_document()
            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
            self._doc.save(output_path)
        finally:
            for prepared in self._temp_images:
                prepared.cleanup()
            self._temp_images.clear()
        return output_path

    def build_answer_key(self) -> Document:
        """Build a separate document with recorded answers."""
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run("Answer Key")
        run.bold = True
        run.font.size = Pt(16)
        doc.add_paragraph(f"{self.fmt.get('subject', '')}  —  Class {self.fmt.get('class', '')}")
        for section in self.fmt.get("sections", []):
            sid = section["id"]
            doc.add_heading(section["title"], level=2)
            for number in range(1, section["total_questions"] + 1):
                if not self.state.is_filled(sid, number):
                    continue
                question = self.state.get_question(sid, number) or {}
                answer = question.get("answer", "")
                if section["question_type"] == "MCQ":
                    options = question.get("options") or []
                    if answer and answer in options:
                        label = OPTION_LABELS[options.index(answer)]
                        answer = f"({label}) {answer}"
                p = doc.add_paragraph()
                printed_number = section.get("start_number", 1) + number - 1
                p.add_run(f"{printed_number}. ").bold = True
                p.add_run(answer or "Answer not entered")
        return doc

    # -- header / instructions -------------------------------------------
    def _configure_base_styles(self) -> None:
        section = self._doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        self._right_tab_cm = 16.8
        style = self._doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

    def _add_header(self) -> None:
        doc = self._doc
        school = self.fmt.get("school_name", "")
        if school:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(school)
            run.bold = True
            run.font.size = Pt(16)

        paper_title = self.fmt.get("paper_title", "")
        if paper_title:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(paper_title)
            run.bold = True
            run.font.size = Pt(14)

        subject = self.fmt.get("subject", "")
        klass = self.fmt.get("class", "")
        title_bits = [b for b in (subject, f"Class {klass}" if klass else "") if b]
        if title_bits:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("  —  ".join(title_bits))
            run.bold = True
            run.font.size = Pt(13)

        # Duration (left) ... Total marks (right) on one tab-balanced line.
        duration = self.fmt.get("duration", "")
        total_marks = self.fmt.get("total_marks", "")
        if duration or total_marks not in ("", None):
            p = doc.add_paragraph()
            self._set_right_tab(p)
            left = f"Time: {duration}" if duration else ""
            right = f"Maximum Marks: {total_marks}" if total_marks not in ("", None) else ""
            run = p.add_run(left)
            run.bold = True
            run = p.add_run("\t" + right)
            run.bold = True

        self._horizontal_rule()

    def _add_instructions(self) -> None:
        instructions = self.fmt.get("instructions") or []
        if not instructions:
            return
        doc = self._doc
        p = doc.add_paragraph()
        run = p.add_run("General Instructions:")
        run.bold = True
        run.underline = True
        try:
            doc.styles["List Number"]
            numbered_style = True
        except KeyError:
            numbered_style = False
        for index, line in enumerate(instructions, start=1):
            item = (
                doc.add_paragraph(str(line), style="List Number")
                if numbered_style else doc.add_paragraph(f"{index}. {line}")
            )
            item.paragraph_format.space_after = Pt(2)
        self._horizontal_rule()

    # -- sections / questions --------------------------------------------
    def _add_section(self, section: dict[str, Any]) -> None:
        doc = self._doc
        sid = section.get("id", "")
        title = section.get("title", f"Section {sid}")

        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(self._template_heading_pt if self.fmt.get("docx_template") else 13)
        if not self.fmt.get("docx_template"):
            self._top_border(heading)

        description = section_instruction(section)
        if description:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            run = p.add_run(str(description))
            run.italic = not bool(self.fmt.get("docx_template"))

        marks_each = section.get("marks_per_question", 0)
        total_q = int(section.get("total_questions", 0))
        for qnum in range(1, total_q + 1):
            question = self.state.get_question(sid, qnum)
            printed_number = section.get("start_number", 1) + qnum - 1
            self._add_question(section, printed_number, question, marks_each)

    def _add_question(
        self,
        section: dict[str, Any],
        qnum: int,
        question: dict[str, Any] | None,
        marks_each: Any,
    ) -> None:
        doc = self._doc
        first_paragraph = len(doc.paragraphs)
        qtype = section.get("question_type", "short_answer")
        text = (question or {}).get("text", "").strip()
        marks = (question or {}).get("marks", marks_each)

        # Main question line with right-aligned marks.
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        self._set_right_tab(p)
        label_run = p.add_run(f"{qnum}. ")
        label_run.bold = True
        body = text if text else "____________________________"
        p.add_run(body)
        if question and question.get("image_path") and question.get("image_position") == "alongside":
            self._add_image(question, p)
        if marks not in ("", None):
            p.add_run(f"\t[{self._fmt_marks(marks)}]")

        if not question:
            return

        if question.get("image_path") and question.get("image_position") == "before_parts":
            self._add_image(question)

        # MCQ options.
        if qtype == "MCQ" and question.get("options"):
            self._add_mcq_options(question["options"], question.get("answer", ""))

        # Sub-parts (short/long answer).
        for label, part in question.get("sub_parts", {}).items():
            self._add_sub_part(label, part)

        # Optional image.
        if question.get("image_path") and question.get("image_position") == "below":
            self._add_image(question)

        # Optional table.
        if question.get("table_def"):
            self._add_table(question["table_def"])
        for paragraph in doc.paragraphs[first_paragraph:-1]:
            paragraph.paragraph_format.keep_with_next = True

    def _add_mcq_options(self, options: list[str], answer: str) -> None:
        doc = self._doc
        for i, opt in enumerate(options):
            if not str(opt).strip():
                continue
            label = OPTION_LABELS[i] if i < len(OPTION_LABELS) else str(i + 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(f"({label}) ")
            run.bold = True
            p.add_run(str(opt))

    def _add_sub_part(self, label: str, part: dict[str, Any]) -> None:
        doc = self._doc
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        self._set_right_tab(p)
        run = p.add_run(f"({label}) ")
        run.bold = True
        p.add_run(part.get("text", "") or "____________________")
        marks = part.get("marks")
        if marks:
            p.add_run(f"\t[{self._fmt_marks(marks)}]")

    def _add_image(self, question: dict[str, Any], paragraph=None) -> None:
        doc = self._doc
        max_width_cm = image_handler.DEFAULT_MAX_WIDTH_CM
        pct = int(question.get("image_width_pct", 100) or 100)
        prepared = image_handler.prepare_image(
            question["image_path"], max_width_cm=max_width_cm
        )
        self._temp_images.append(prepared)
        width_cm = prepared.width_cm * max(10, min(100, pct)) / 100.0
        p = paragraph or doc.add_paragraph()
        if paragraph is None:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(prepared.embed_path, width=Cm(width_cm))

    def _add_table(self, table_def: dict[str, Any]) -> None:
        rows = int(table_def.get("rows", 1))
        cols = int(table_def.get("cols", 1))
        headers = table_def.get("headers") or None
        table_builder.build_table(self._doc, rows, cols, headers)
        self._doc.add_paragraph()  # spacing after the table

    # -- low-level formatting helpers ------------------------------------
    def _set_right_tab(self, paragraph) -> None:
        """Add a right-aligned tab stop so marks sit flush against the margin."""
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(self._right_tab_cm), WD_TAB_ALIGNMENT.RIGHT
        )

    def _horizontal_rule(self) -> None:
        """A thin full-width rule, drawn as a paragraph bottom border."""
        p = self._doc.add_paragraph()
        self._set_border(p, "bottom")

    def _top_border(self, paragraph) -> None:
        self._set_border(paragraph, "top")

    @staticmethod
    def _set_border(paragraph, edge: str) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "1")
        element.set(qn("w:color"), "808080")
        borders.append(element)

    @staticmethod
    def _fmt_marks(marks: Any) -> str:
        try:
            f = float(marks)
        except (TypeError, ValueError):
            return str(marks)
        return str(int(f)) if f.is_integer() else f"{f:g}"
