"""Helpers for building tables inside the generated document."""

from __future__ import annotations

from typing import Optional

from docx.enum.table import WD_TABLE_ALIGNMENT


def build_table(doc, rows: int, cols: int, headers: Optional[list[str]] = None):
    """Append a ``rows`` x ``cols`` table to ``doc`` and return it.

    If ``headers`` is given, the first row is filled with the header labels and
    rendered bold. The table is left-aligned with a single-line grid style.
    """
    rows = max(1, int(rows))
    cols = max(1, int(cols))

    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    try:
        table.style = "Table Grid"
    except KeyError:
        # Style not in the default template; the table still renders.
        pass

    if headers:
        header_cells = table.rows[0].cells
        for i in range(cols):
            label = headers[i] if i < len(headers) else ""
            cell = header_cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(label))
            run.bold = True

    return table
