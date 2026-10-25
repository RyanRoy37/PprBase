"""Import a Word paper as a reusable layout, keeping its original DOCX."""

from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
from copy import deepcopy
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

from core.format_loader import validate_format


class WordTemplateError(Exception):
    """The selected Word file cannot be used as a template."""


GROUP_MARKS = re.compile(r"\(?\s*(\d+)\s*[x×]\s*(\d+(?:\.\d+)?|\d+/\d+)\s*=\s*(\d+(?:\.\d+)?)\s*\)?", re.I)
SECTION = re.compile(r"^\s*(?:section|part)\s*[-: ]*\s*(?:[ivx]+|\d+|[a-z])\b", re.I)
QUESTION = re.compile(r"^\s*question\s*[-: ]*\s*(\d+)\b", re.I)
ANSWER_ANY = re.compile(r"answer\s+any\s+(\d+)", re.I)
MARKS = re.compile(r"\bmarks?\s*[:\-]\s*(\d+(?:\.\d+)?)", re.I)
SUBJECT = re.compile(r"\bsubject\s*[:\-]\s*(.+)", re.I)
CLASS = re.compile(r"\bclass\s*[:\-]\s*([a-z0-9]+)", re.I)
TIME = re.compile(r"\b(?:time|duration)\s*[:\-]\s*(.+)", re.I)


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _number(value: str) -> float:
    if "/" in value:
        numerator, denominator = value.split("/", 1)
        return int(numerator) / int(denominator)
    return float(value)


def _question_type(title: str) -> str:
    title = title.lower()
    if any(word in title for word in ("choose", "multiple choice", "assertion", "case study")):
        return "MCQ"
    if any(word in title for word in ("essay", "long answer", "explain in detail")):
        return "long_answer"
    return "short_answer"


def _body_paragraphs(doc) -> list[tuple[int, str]]:
    result = []
    for index, element in enumerate(doc.element.body.iterchildren()):
        if element.tag == qn("w:p"):
            text = _clean(Paragraph(element, doc).text)
            if text:
                result.append((index, text))
    return result


def template_body_choices(path: str | Path) -> list[tuple[int, str]]:
    """Readable body rows for choosing where old questions are replaced."""
    doc = Document(path)
    choices = []
    for index, element in enumerate(doc.element.body.iterchildren()):
        if element.tag == qn("w:p"):
            text = _clean(Paragraph(element, doc).text)
        elif element.tag == qn("w:tbl"):
            table = Table(element, doc)
            text = "Table: " + _clean(" ".join(cell.text for row in table.rows[:2] for cell in row.cells))
        else:
            continue
        if text:
            choices.append((index, text[:100]))
    return choices


def inspect_word_template(path: str | Path) -> dict:
    """Propose an editable format from headings and mark formulas in a DOCX.

    The teacher reviews this proposal before it is saved. The stored DOCX
    supplies the header, page setup, and base styles for future exports.
    """
    source = Path(path)
    if source.suffix.lower() != ".docx" or not source.is_file():
        raise WordTemplateError("Choose an existing .docx Word paper.")
    try:
        doc = Document(source)
    except Exception as exc:
        raise WordTemplateError(f"Couldn't read the Word paper: {exc}") from exc
    records = _body_paragraphs(doc)
    if not records:
        raise WordTemplateError("The Word paper has no readable paragraphs.")

    first_section = next((index for index, text in records if SECTION.match(text)), None)
    first_group = next((index for index, text in records if GROUP_MARKS.search(text)), None)
    first_question = next((index for index, text in records if QUESTION.match(text)), None)
    body_start = first_section if first_section is not None else (
        first_group if first_group is not None else first_question
    )
    if body_start is None:
        # Use the first body row after the common school/title/subject/class
        # header. The review screen lets the teacher correct this boundary.
        body_start = records[min(4, len(records) - 1)][0]

    header = [text for index, text in records if index < body_start]
    school = header[0] if header else "School name"
    title = next((text for text in header[1:] if re.search(r"\b(?:test|exam|assessment|paper)\b", text, re.I)), "")
    subject = next((match.group(1).strip() for text in header if (match := SUBJECT.search(text))), "Subject")
    klass = next((match.group(1).strip() for text in header if (match := CLASS.search(text))), "Class")
    maximum = next((float(match.group(1)) for text in header if (match := MARKS.search(text))), None)
    duration = next((match.group(1).strip() for text in header if (match := TIME.search(text))), "")

    sections = []
    major = ""
    trailing_questions: list[int] = []
    trailing_choice = None
    trailing_title = ""
    for index, text in records:
        if index < body_start:
            continue
        if SECTION.match(text):
            major = text
            trailing_title = text
            trailing_questions = []
            trailing_choice = None
        choice = ANSWER_ANY.search(text)
        if choice:
            trailing_choice = int(choice.group(1))
        group = GROUP_MARKS.search(text)
        if group:
            trailing_questions = []
            count = int(group.group(1))
            marks_each = _number(group.group(2))
            heading = _clean(GROUP_MARKS.sub("", text)).strip(" .:-")
            heading = re.sub(r"^[a-z]\.\s*", "", heading, flags=re.I)
            if not heading:
                heading = f"Questions worth {marks_each:g} marks"
            if not sections and major:
                heading = f"{major} — {heading}"
            sid = chr(ord("A") + len(sections)) if len(sections) < 26 else f"S{len(sections) + 1}"
            sections.append({
                "id": sid, "title": heading, "description": "",
                "question_type": _question_type(heading),
                "total_questions": count, "marks_per_question": marks_each,
                "sub_parts": None,
            })
        question = QUESTION.match(text)
        if question and trailing_title and not GROUP_MARKS.search(trailing_title):
            trailing_questions.append(int(question.group(1)))

    # Questions after the final section heading form a choice section in the
    # supplied examples. Only the required answers count toward maximum marks.
    if trailing_questions:
        printed = len(dict.fromkeys(trailing_questions))
        required = min(trailing_choice or printed, printed)
        already = sum(s["total_questions"] * s["marks_per_question"] for s in sections)
        remaining = max(0.0, (maximum or already + printed) - already)
        marks_each = remaining / required if required else 1.0
        sid = chr(ord("A") + len(sections)) if len(sections) < 26 else f"S{len(sections) + 1}"
        sections.append({
            "id": sid, "title": trailing_title or "Questions", "description": "",
            "question_type": "short_answer", "total_questions": printed,
            "attempt_count": required, "start_number": min(trailing_questions),
            "marks_per_question": marks_each, "sub_parts": None,
        })
    if not sections:
        numbered = [int(match.group(1)) for index, text in records if index >= body_start if (match := QUESTION.match(text))]
        count = len(set(numbered)) or 1
        sections = [{
            "id": "A", "title": "Questions", "description": "",
            "question_type": "short_answer", "total_questions": count,
            "marks_per_question": (maximum or count) / count,
            "sub_parts": None,
        }]
    if maximum is None:
        maximum = sum(s.get("attempt_count", s["total_questions"]) * s["marks_per_question"] for s in sections)

    fmt = {
        "format_id": re.sub(r"[^a-z0-9]+", "_", f"{source.stem}_{subject}_{klass}".lower()).strip("_")[:72],
        "display_name": f"{title or source.stem} — {subject} Class {klass}",
        "school_name": school,
        "paper_title": title,
        "subject": subject,
        "class": klass,
        "total_marks": maximum,
        "duration": duration,
        "instructions": [],
        "sections": sections,
        "docx_template": str(source.resolve()),
        "template_body_start": body_start,
    }
    validate_format(fmt)
    return fmt


def install_word_template(fmt: dict, source_path: str | Path, formats_dir: str | Path) -> str:
    """Store the original DOCX beside formats and create a reusable format card."""
    source = Path(source_path)
    if not source.is_file():
        raise WordTemplateError("The selected Word paper has moved. Choose it again.")
    data = deepcopy(fmt)
    validate_format(data)
    directory = Path(formats_dir)
    directory.mkdir(parents=True, exist_ok=True)
    templates_dir = directory.parent / "templates"
    templates_dir.mkdir(parents=True, exist_ok=True)
    stem = data["format_id"]
    while (directory / f"{stem}.json").exists() or (templates_dir / f"{stem}.docx").exists():
        stem = f"{data['format_id']}_{uuid4().hex[:6]}"
    data["format_id"] = stem
    target = templates_dir / f"{stem}.docx"
    data["docx_template"] = os.path.relpath(target, directory)
    format_path = directory / f"{stem}.json"
    try:
        shutil.copy2(source, target)
        with format_path.open("x", encoding="utf-8") as file:
            json.dump(data, file, ensure_ascii=False, indent=2)
    except Exception:
        target.unlink(missing_ok=True)
        format_path.unlink(missing_ok=True)
        raise
    return str(format_path)


def revise_word_template(fmt: dict, source_path: str | Path, format_path: str | Path) -> str:
    """Save a new Word source for an existing format without changing its card ID.

    Each revision gets its own file, so existing drafts and duplicated format
    cards can continue using the Word source they were created with.
    """
    source = Path(source_path)
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise WordTemplateError("Choose an existing .docx Word paper.")
    try:
        Document(source)
    except Exception as exc:
        raise WordTemplateError(f"Couldn't read the replacement Word paper: {exc}") from exc
    data = deepcopy(fmt)
    validate_format(data)
    path = Path(format_path)
    if not path.is_file():
        raise WordTemplateError("The saved template card could not be found.")
    templates_dir = path.parent.parent / "templates"
    templates_dir.mkdir(parents=True, exist_ok=True)
    target = templates_dir / f"{data['format_id']}_{uuid4().hex[:8]}.docx"
    data["docx_template"] = os.path.relpath(target, path.parent)
    temporary = None
    try:
        shutil.copy2(source, target)
        fd, temporary = tempfile.mkstemp(prefix=".format_", suffix=".json", dir=path.parent)
        with os.fdopen(fd, "w", encoding="utf-8") as file:
            json.dump(data, file, ensure_ascii=False, indent=2)
            file.flush()
            os.fsync(file.fileno())
        os.replace(temporary, path)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
    return str(path)
