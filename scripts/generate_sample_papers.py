"""Generate the three complete Word, PDF, and answer-key examples in samples/."""

from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from PyQt6.QtPdf import QPdfDocument
from PyQt6.QtWidgets import QApplication

from core.document_engine import DocumentEngine
from core.pdf_engine import PdfEngine
from tests.scenarios import build_scenario


DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "samples"


def generate(output_dir: Path) -> list[dict]:
    app = QApplication.instance() or QApplication([])
    output_dir.mkdir(parents=True, exist_ok=True)
    report = []
    with tempfile.TemporaryDirectory(prefix="pprbase_samples_") as temporary:
        for marks in (25, 50, 100):
            fmt, state = build_scenario(marks, Path(temporary) / str(marks))
            stem = f"{marks}_mark_{fmt['subject'].lower()}"
            paper_path = output_dir / f"{stem}.docx"
            pdf_path = output_dir / f"{stem}.pdf"
            key_path = output_dir / f"{stem}_answers.docx"
            engine = DocumentEngine(fmt, state)
            engine.save(str(paper_path))
            engine.build_answer_key().save(str(key_path))
            PdfEngine(fmt, state).save(str(pdf_path))

            paper = Document(paper_path)
            key = Document(key_path)
            expected_questions = sum(section["total_questions"] for section in fmt["sections"])
            paper_questions = [p for p in paper.paragraphs if re.match(r"^\d+\. ", p.text)]
            key_questions = [p for p in key.paragraphs if re.match(r"^\d+\. ", p.text)]
            if state.allocated_marks(fmt) != marks or len(paper_questions) != expected_questions or len(key_questions) != expected_questions:
                raise RuntimeError(f"The {marks}-mark sample failed validation.")
            for path in (paper_path, key_path):
                with zipfile.ZipFile(path) as archive:
                    if archive.testzip() is not None:
                        raise RuntimeError(f"The saved document is damaged: {path}")
            pdf = QPdfDocument(None)
            if pdf.load(str(pdf_path)) != QPdfDocument.Error.None_ or pdf.pageCount() < 1:
                raise RuntimeError(f"The saved PDF is damaged: {pdf_path}")
            pdf_pages = pdf.pageCount()
            pdf.close()
            report.append({
                "marks": marks,
                "subject": fmt["subject"],
                "questions": expected_questions,
                "paper": paper_path.name,
                "pdf": pdf_path.name,
                "pdf_pages": pdf_pages,
                "answer_key": key_path.name,
                "embedded_images": len(paper.inline_shapes),
                "tables": len(paper.tables),
                "status": "passed",
            })
    (output_dir / "test_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def main() -> None:
    if os.name != "nt" and not os.environ.get("DISPLAY") and not os.environ.get("WAYLAND_DISPLAY"):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    for item in generate(args.output_dir):
        print(f"{item['marks']} marks: {item['questions']} questions, {item['embedded_images']} image(s), {item['tables']} table(s) — passed")


if __name__ == "__main__":
    main()
