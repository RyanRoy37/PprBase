"""Load the two supplied 2024 Chemistry papers as editable PprBase drafts.

This is an explicit transcription of these two documents, not a general Word
importer. The source hashes protect the paragraph mapping from silently being
used on a different revision of either file.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from PyQt6.QtPdf import QPdfDocument
from PyQt6.QtWidgets import QApplication

from core.document_engine import DocumentEngine
from core.draft_store import DraftStore
from core.format_loader import calculated_total_marks, validate_format
from core.pdf_engine import PdfEngine
from core.state_manager import PaperState

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "VI": (ROOT / "tests" / "fixtures" / "chemistry" / "Chemistry Ut1 VI Blueprint.docx", "b86887e77a9e6943977a28acd37159a56fd77d8bebef2dc12c448df103e00de1"),
    "VII": (ROOT / "tests" / "fixtures" / "chemistry" / "Chemistry Ut1 VII qs paper 2024.docx", "0f490e4eb655ed9d3ae5f5dd1588e67c3b683817b8948eca5b7f64491e227156"),
}
PART_MARKS = (1, 1, 0.5, 1, 1, 0.5)
PART_LABELS = ("i", "ii", "iii", "iv", "v", "vi")


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().strip(",")


def _without_mark(value: str) -> str:
    return re.sub(r"\s*\((?:1|1/2)\)\s*\.?$", "", _clean(value)).strip()


def _part_text(value: str) -> str:
    value = re.sub(r"\s*\((?:1|1/2)\)\.?(?=\s|$)", " ", _clean(value))
    value = re.sub(r"^(?:(?:[a-c]\.\s*)?[ivx]+\.\s*)", "", value, flags=re.I)
    return _clean(value)


def _section(sid: str, title: str, qtype: str, count: int, marks: float, description: str = "") -> dict:
    return {
        "id": sid, "title": title, "description": description,
        "question_type": qtype, "total_questions": count,
        "marks_per_question": marks, "sub_parts": None,
    }


def _base_format(klass: str) -> dict:
    fmt = {
        "format_id": f"chemistry_ut1_class_{klass.lower()}",
        "display_name": f"Unit Test 1 — Chemistry Class {klass}",
        "school_name": "St Vincent Pallotti School",
        "paper_title": "Unit Test-1",
        "subject": "Chemistry",
        "class": klass,
        "total_marks": 25,
        "duration": "",
        "instructions": [],
        "sections": [
            _section("A", "Section I — Question 1(a): Multiple Choice", "MCQ", 6, 0.5, "Choose the correct answer."),
            _section("B", "Question 1(b): Short Answer", "short_answer", 3, 1),
            _section("C", "Question 1(c): Assertion and Reason", "MCQ", 3, 1),
            _section("D", "Question 1(d): Chemistry", "short_answer", 3, 1),
            _section("E", "Question 1(e): Case Study", "MCQ", 3, 1, "Read the passage and answer the questions."),
            _section("F", "Section II — Questions 2 to 4", "short_answer", 3, 5, "Answer any 2 questions."),
        ],
        "docx_template": None,
    }
    fmt["sections"][5].update({"attempt_count": 2, "start_number": 2, "sub_parts": list(PART_LABELS)})
    return fmt


def _add(state: PaperState, sid: str, number: int, text: str, marks: float, *, options: list[str] | None = None,
         image: str | None = None, parts: list[str] | None = None) -> None:
    data = {"text": _without_mark(text), "marks": marks, "type": "MCQ" if options is not None else "short_answer"}
    if options is not None:
        data["options"] = [_clean(option) for option in options]
    if image:
        data.update({"image_path": image, "image_position": "before_parts", "image_width_pct": 65})
    if parts is not None:
        if len(parts) != 6:
            raise ValueError("Each Section II question must have six parts.")
        data["sub_parts"] = {
            label: {"text": _part_text(part), "marks": mark}
            for label, part, mark in zip(PART_LABELS, parts, PART_MARKS)
        }
    state.add_question(sid, number, data)


def _asset(source: Path, internal_name: str, asset_dir: Path, target_name: str) -> str:
    asset_dir.mkdir(parents=True, exist_ok=True)
    target = asset_dir / target_name
    with ZipFile(source) as archive:
        target.write_bytes(archive.read(f"word/media/{internal_name}"))
    return str(target)


def load_reference(klass: str, source: Path, asset_dir: Path) -> tuple[dict, PaperState]:
    expected_hash = SOURCES[klass][1]
    actual_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    if actual_hash != expected_hash:
        raise ValueError(f"{source.name} differs from the supplied version; review its paragraph mapping before importing.")
    paragraphs = [p.text for p in Document(source).paragraphs]
    fmt = _base_format(klass)
    state = PaperState(fmt["format_id"])

    if klass == "VII":
        atom = _asset(source, "image1.png", asset_dir, "class_vii_atom.png")
        funnel = _asset(source, "image2.png", asset_dir, "class_vii_funnel.png")
        sublimation = _asset(source, "image6.png", asset_dir, "class_vii_sublimation.png")
        mcq = [
            (7, ["Evaporation", "Distillation", "Sublimation", "Fractional Distillation"]),
            (9, ["1 & 2", "2 & 3", "3 & 4", "2 & 4"]),
            (11, ["Boron", "Radon", "Sulphur", "Hydrogen"]),
            (13, ["Carbonate", "Bicarbonate", "Ammonium", "Hydroxide"]),
            (15, ["Chalk & water", "Iodine & alcohol", "Kerosene & water", "Iron & sulphur"]),
            (17, ["Silver", "Gold", "Potassium", "Iron"]),
        ]
        for number, (index, options) in enumerate(mcq, 1):
            _add(state, "A", number, re.sub(r"^\s*\d+\.\s*", "", paragraphs[index]), 0.5, options=options)
        for number, index in enumerate((21, 22, 23), 1):
            _add(state, "B", number, paragraphs[index], 1)
        for number, (stem, reason, choices) in enumerate(((26, 27, (28, 29, 30, 31)), (32, 33, (34, 35, 36, 37)), (38, 39, (40, 41, 42, 43))), 1):
            _add(state, "C", number, f"{_clean(paragraphs[stem])}\n{_clean(paragraphs[reason])}", 1,
                 options=[re.sub(r"^(?:i|ii|iii|iv)\.\s*", "", _clean(paragraphs[i]), flags=re.I) for i in choices])
        for number, index in enumerate((46, 47, 48), 1):
            _add(state, "D", number, paragraphs[index], 1)
        _add(state, "E", 1, f"{_clean(paragraphs[50])}\n\n{_clean(paragraphs[51])}", 1,
             options=["5", "6", "7", "8"], image=atom)
        _add(state, "E", 2, paragraphs[53], 1,
             options=["Moving around the nucleus", "In the electron shells", "In the energy levels", "In the nucleus"])
        _add(state, "E", 3, paragraphs[56], 1,
             options=["The number of neutrons", "The number of protons and electrons", "The number of electrons", "The sum of protons and neutrons"])
        parts = [
            [paragraphs[64], paragraphs[65], paragraphs[66] + " " + paragraphs[67], paragraphs[69],
             paragraphs[70] + " " + paragraphs[71] + " " + paragraphs[72] + " " + paragraphs[73], paragraphs[74]],
            [paragraphs[77] + " " + paragraphs[78], paragraphs[79] + " " + paragraphs[80],
             paragraphs[81] + " " + paragraphs[82], paragraphs[84], paragraphs[85], paragraphs[86]],
            [paragraphs[89] + " " + paragraphs[90] + " " + paragraphs[91], paragraphs[92], paragraphs[93],
             paragraphs[94], paragraphs[95], paragraphs[96] + " " + paragraphs[97]],
        ]
        for number, items in enumerate(parts, 1):
            _add(state, "F", number, "Answer the following parts.", 5,
                 image={1: funnel, 3: sublimation}.get(number), parts=items)
    else:
        steel = _asset(source, "image1.png", asset_dir, "class_vi_steel_pipes.png")
        mcq = [
            (7, ["Elements", "Compounds", "Mixture", "Metalloid"]),
            (10, ["Ca OH", "Ca2 OH", "CA (OH)2", "Ca (OH)2"]),
            (12, ["HCl", "NH Cl", "CH3Cl", "H2O2"]),
            (14, ["Zinken", "Natrium", "Cuprum", "Aurum"]),
            (16, ["NH3", "HNO3", "KNO3", "KOH"]),
            (19, ["Nitrogen", "Chlorine", "Helium", "Hydrogen"]),
        ]
        for number, (index, options) in enumerate(mcq, 1):
            _add(state, "A", number, paragraphs[index], 0.5, options=options)
        for number, index in enumerate((24, 25, 26), 1):
            _add(state, "B", number, paragraphs[index], 1)
        for number, (stem, reason, choices) in enumerate(((29, 30, (31, 32, 33, 34)), (35, 36, (37, 38, 39, 40)), (41, 42, (43, 44, 45, 46))), 1):
            _add(state, "C", number, f"{_clean(paragraphs[stem])}\n{_clean(paragraphs[reason])}", 1,
                 options=[paragraphs[i] for i in choices])
        for number, index in enumerate((50, 51, 52), 1):
            _add(state, "D", number, paragraphs[index], 1)
        _add(state, "E", 1, f"{_clean(paragraphs[56])}\n\n{_clean(paragraphs[57])}", 1, options=[])
        _add(state, "E", 2, paragraphs[58], 1, options=["Mass number", "Atomic mass", "Atomic number", "Valency"])
        _add(state, "E", 3, paragraphs[60], 1, options=["Nitrogen", "Hydrogen", "Oxygen", "Helium"])
        parts = [
            [paragraphs[68], paragraphs[69], paragraphs[70] + " " + paragraphs[71], paragraphs[72] + " " + paragraphs[73],
             paragraphs[74] + " " + paragraphs[75], paragraphs[76]],
            [paragraphs[82] + " " + paragraphs[83], paragraphs[84] + " " + paragraphs[85], paragraphs[86],
             paragraphs[88] + " " + paragraphs[89], paragraphs[90] + " [Zinc, Sulphur, Phosphorus, Mercury, Carbon, Ferrum]", paragraphs[92]],
            [paragraphs[95], paragraphs[96], paragraphs[97], paragraphs[99], paragraphs[100], paragraphs[101]],
        ]
        for number, items in enumerate(parts, 1):
            _add(state, "F", number, "Answer the following parts.", 5,
                 image=steel if number == 3 else None, parts=items)

    validate_format(fmt)
    if calculated_total_marks(fmt) != 25 or state.allocated_marks(fmt) != 25:
        raise ValueError(f"Class {klass} did not produce a complete 25-mark paper.")
    if any(state.count_filled(section["id"]) != section["total_questions"] for section in fmt["sections"]):
        raise ValueError(f"Class {klass} is missing a question.")
    return fmt, state


def import_and_test(source_paths: dict[str, Path], root: Path = ROOT, refresh: bool = False) -> list[dict]:
    app = QApplication.instance() or QApplication([])
    asset_dir = root / "samples" / "reference_assets"
    sample_dir = root / "samples" / "reference_papers"
    sample_dir.mkdir(parents=True, exist_ok=True)
    report = []
    for klass in ("VI", "VII"):
        source = source_paths[klass]
        fmt, state = load_reference(klass, source, asset_dir)
        format_path = sample_dir / "legacy_formats" / f"{fmt['format_id']}.json"
        if not format_path.exists():
            format_path.parent.mkdir(parents=True, exist_ok=True)
            format_path.write_text(json.dumps(fmt, ensure_ascii=False, indent=2), encoding="utf-8")
        draft_id = f"chemistryut1{klass.lower()}{hashlib.sha256(source.read_bytes()).hexdigest()[:8]}"
        draft_store = DraftStore(str(root / "drafts"))
        if refresh or not (draft_store.directory / f"{draft_id}.json").exists():
            draft_store.save(fmt, state, draft_id)
        restored_fmt, restored = draft_store.load(draft_id)
        stem = f"chemistry_ut1_class_{klass.lower()}"
        docx_path = sample_dir / f"{stem}.docx"
        pdf_path = sample_dir / f"{stem}.pdf"
        DocumentEngine(restored_fmt, restored).save(str(docx_path))
        PdfEngine(restored_fmt, restored).save(str(pdf_path))
        pdf = QPdfDocument(None)
        if pdf.load(str(pdf_path)) != QPdfDocument.Error.None_:
            raise ValueError(f"The Class {klass} PDF could not be opened.")
        pages = pdf.pageCount()
        pdf.close()
        report.append({
            "class": klass, "source": source.name, "format": format_path.name,
            "draft_id": draft_id, "questions": sum(s["total_questions"] for s in fmt["sections"]),
            "maximum_marks": 25, "printed_choice_questions": 3, "questions_to_answer": 2,
            "docx": docx_path.name, "pdf": pdf_path.name, "pdf_pages": pages,
        })
    (sample_dir / "import_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--class-vi", type=Path, default=SOURCES["VI"][0])
    parser.add_argument("--class-vii", type=Path, default=SOURCES["VII"][0])
    parser.add_argument("--refresh", action="store_true", help="Replace these imported drafts with the source content")
    args = parser.parse_args()
    if os.name != "nt" and not os.environ.get("DISPLAY") and not os.environ.get("WAYLAND_DISPLAY"):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    for item in import_and_test({"VI": args.class_vi, "VII": args.class_vii}, refresh=args.refresh):
        print(f"Class {item['class']}: {item['questions']} question slots, 25 marks, {item['pdf_pages']} PDF pages — imported")


if __name__ == "__main__":
    main()
