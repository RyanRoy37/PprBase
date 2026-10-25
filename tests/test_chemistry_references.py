"""Regression checks for the two supplied Class VI and VII papers."""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PyQt6.QtPdf import QPdfDocument
from PyQt6.QtWidgets import QApplication

from core.document_engine import DocumentEngine
from core.draft_store import DraftStore
from core.format_loader import calculated_total_marks, section_instruction
from core.pdf_engine import PdfEngine
from scripts.import_chemistry_references import SOURCES, load_reference
from ui.main_window import MainWindow


class ChemistryReferenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def test_both_supplied_papers_import_and_export(self) -> None:
        for klass, expected_images in (("VI", 1), ("VII", 3)):
            with self.subTest(klass=klass), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                fmt, state = load_reference(klass, SOURCES[klass][0], root / "images")
                self.assertEqual(fmt["paper_title"], "Unit Test-1")
                self.assertEqual(calculated_total_marks(fmt), 25)
                self.assertEqual(state.allocated_marks(fmt), 25)
                self.assertEqual([s["total_questions"] for s in fmt["sections"]], [6, 3, 3, 3, 3, 3])
                self.assertEqual(fmt["sections"][-1]["attempt_count"], 2)
                self.assertEqual(fmt["sections"][-1]["start_number"], 2)
                self.assertEqual(section_instruction(fmt["sections"][-1]), "Answer any 2 questions.")
                self.assertTrue(all(state.count_filled(s["id"]) == s["total_questions"] for s in fmt["sections"]))
                self.assertEqual(
                    [sum(p["marks"] for p in state.get_question("F", n)["sub_parts"].values()) for n in (1, 2, 3)],
                    [5, 5, 5],
                )

                store = DraftStore(str(root / "drafts"))
                restored_fmt, restored = store.load(store.save(fmt, state))
                docx_path = root / "paper.docx"
                pdf_path = root / "paper.pdf"
                DocumentEngine(restored_fmt, restored).save(str(docx_path))
                PdfEngine(restored_fmt, restored).save(str(pdf_path))
                paper = Document(docx_path)
                paper_text = "\n".join(p.text for p in paper.paragraphs)
                self.assertIn("St Vincent Pallotti School", paper_text)
                self.assertIn("Unit Test-1", paper_text)
                self.assertIn("Maximum Marks: 25", paper_text)
                self.assertIn("Answer any 2 questions.", paper_text)
                self.assertEqual(len(paper.inline_shapes), expected_images)
                self.assertEqual(
                    [p.text.split(".", 1)[0] for p in paper.paragraphs if p.text.startswith(("2. Answer", "3. Answer", "4. Answer"))],
                    ["2", "3", "4"],
                )
                self.assertIn("[0.5]", paper_text)

                pdf = QPdfDocument(None)
                self.assertEqual(pdf.load(str(pdf_path)), QPdfDocument.Error.None_)
                self.assertGreaterEqual(pdf.pageCount(), 1)
                page_text = PdfEngine._flat("\n".join(pdf.getAllText(i).text() for i in range(pdf.pageCount())))
                self.assertIn("Maximum Marks: 25", page_text)
                self.assertIn("Answer any 2 questions.", page_text)
                self.assertIn("4. Answer the following parts.", page_text)
                pdf.close()
                self.assertEqual(PdfEngine(restored_fmt, restored)._orphaned_sections(pdf_path), set())
                self.assertEqual(PdfEngine(restored_fmt, restored)._orphaned_questions(pdf_path), set())

                window = MainWindow({"paths": {"formats_dir": str(root / "formats"), "drafts_dir": str(root / "drafts"), "output_dir": str(root)}})
                window._show_editor(restored_fmt, restored)
                with patch("ui.main_window.QMessageBox.question", side_effect=AssertionError("Complete paper was flagged")):
                    self.assertTrue(window._confirm_incomplete())
                window.close()

    def test_choice_marks_count_only_required_answers(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            fmt, state = load_reference("VI", SOURCES["VI"][0], Path(temporary) / "images")
            self.assertEqual(state.allocated_marks(fmt), 25)
            state.delete_question("F", 3)
            self.assertEqual(state.allocated_marks(fmt), 25)
            state.delete_question("F", 2)
            self.assertEqual(state.allocated_marks(fmt), 20)


if __name__ == "__main__":
    unittest.main()
