"""Full 25, 50 and 100 mark paper export checks."""

from __future__ import annotations

import os
import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Cm
from PyQt6.QtCore import QSize
from PyQt6.QtPdf import QPdfDocument
from PyQt6.QtWidgets import QApplication

from core.draft_store import DraftStore
from core.document_engine import DocumentEngine
from core.image_handler import ImageError
from core.pdf_engine import PdfEngine
from tests.scenarios import build_scenario
from ui.main_window import MainWindow


class CompletePaperTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def test_25_50_and_100_mark_papers(self) -> None:
        for total_marks in (25, 50, 100):
            with self.subTest(total_marks=total_marks), tempfile.TemporaryDirectory() as temporary:
                self._check_paper(total_marks, Path(temporary))

    def _check_paper(self, total_marks: int, root: Path) -> None:
        fmt, state = build_scenario(total_marks, root / "source_images")
        question_count = sum(section["total_questions"] for section in fmt["sections"])
        self.assertEqual(state.allocated_marks(fmt), total_marks)
        for section in fmt["sections"]:
            self.assertEqual(state.count_filled(section["id"]), section["total_questions"])
            for number in range(1, section["total_questions"] + 1):
                question = state.get_question(section["id"], number)
                self.assertTrue(question["text"])
                self.assertTrue(question["answer"])
                self.assertEqual(question["marks"], section["marks_per_question"])
                if section["question_type"] == "MCQ":
                    self.assertEqual(len(question["options"]), 4)
                    self.assertIn(question["answer"], question["options"])

        store = DraftStore(str(root / "drafts"))
        draft_id = store.save(fmt, state)
        for image in (root / "source_images").iterdir():
            image.unlink()
        restored_fmt, restored = store.load(draft_id)
        self.assertEqual(restored_fmt, fmt)
        self.assertEqual(restored.allocated_marks(restored_fmt), total_marks)
        for section in fmt["sections"]:
            for number in range(1, section["total_questions"] + 1):
                before = state.get_question(section["id"], number)
                after = restored.get_question(section["id"], number)
                for field in ("text", "answer", "marks", "options", "sub_parts", "table_def"):
                    self.assertEqual(after[field], before[field])
                if before["image_path"]:
                    self.assertTrue(Path(after["image_path"]).is_file())

        paper_path = root / f"{total_marks}_paper.docx"
        key_path = root / f"{total_marks}_answer_key.docx"
        pdf_path = root / f"{total_marks}_paper.pdf"
        window = MainWindow({"paths": {
            "formats_dir": str(root / "formats"),
            "drafts_dir": str(root / "drafts"),
            "output_dir": str(root),
        }, "editor": {"auto_save_minutes": 2}})
        window.show()
        window._show_editor(restored_fmt, restored)
        with patch("ui.main_window.QFileDialog.getSaveFileName", side_effect=[
            (str(paper_path), ""), (str(key_path), ""), (str(pdf_path), ""),
        ]), patch.object(window, "_show_success"), patch(
            "ui.main_window.QMessageBox.question", side_effect=AssertionError("A complete paper was flagged incomplete")
        ):
            window._export()
            window._export_answer_key()
            window._save_pdf()
        with patch("ui.main_window.PdfPreviewDialog.exec", return_value=0) as show_preview, patch(
            "ui.main_window.QMessageBox.warning", side_effect=AssertionError("Preview failed")
        ):
            window._quick_preview()
            show_preview.assert_called_once()
        window.close()

        for path in (paper_path, key_path):
            self.assertTrue(path.is_file())
            self.assertGreater(path.stat().st_size, 20_000)
            with zipfile.ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())

        self.assertTrue(pdf_path.is_file())
        self.assertEqual(pdf_path.read_bytes()[:5], b"%PDF-")
        pdf = QPdfDocument(None)
        self.assertEqual(pdf.load(str(pdf_path)), QPdfDocument.Error.None_)
        self.assertGreaterEqual(pdf.pageCount(), 1)
        page_size = pdf.pagePointSize(0)
        self.assertAlmostEqual(page_size.width(), 595.28, delta=1)
        self.assertAlmostEqual(page_size.height(), 841.89, delta=1)
        pdf_text = PdfEngine._flat("\n".join(pdf.getAllText(page).text() for page in range(pdf.pageCount())))
        self.assertIn(f"Maximum Marks: {total_marks}", pdf_text)
        self.assertIn(fmt["school_name"], pdf_text)
        for section in fmt["sections"]:
            self.assertIn(section["title"], pdf_text)
            for number in range(1, section["total_questions"] + 1):
                words = PdfEngine._flat(restored.get_question(section["id"], number)["text"]).split()
                self.assertIn(" ".join(words[:5]), pdf_text)
                self.assertIn(" ".join(words[-5:]), pdf_text)
        self.assertFalse(pdf.render(0, QSize(420, 594)).isNull())
        pdf.close()
        self.assertEqual(PdfEngine(restored_fmt, restored)._orphaned_sections(pdf_path), set())

        paper = Document(paper_path)
        key = Document(key_path)
        paragraphs = paper.paragraphs
        self.assertAlmostEqual(paper.sections[0].page_width.cm, Cm(21).cm, places=2)
        self.assertAlmostEqual(paper.sections[0].page_height.cm, Cm(29.7).cm, places=2)
        self.assertAlmostEqual(paper.sections[0].left_margin.cm, 2.0, places=2)
        self.assertAlmostEqual(paper.sections[0].right_margin.cm, 2.0, places=2)
        self.assertEqual(len(paper.tables), {25: 1, 50: 1, 100: 2}[total_marks])
        self.assertEqual(len(paper.inline_shapes), {25: 1, 50: 1, 100: 2}[total_marks])
        expected_tables = [
            question["table_def"]
            for section in fmt["sections"]
            for number in range(1, section["total_questions"] + 1)
            if (question := restored.get_question(section["id"], number))["table_def"]
        ]
        for actual, expected in zip(paper.tables, expected_tables):
            self.assertEqual(len(actual.rows), expected["rows"])
            self.assertEqual(len(actual.columns), expected["cols"])
            self.assertEqual([cell.text for cell in actual.rows[0].cells], expected["headers"])

        question_paragraphs = [p for p in paragraphs if re.match(r"^\d+\. ", p.text)]
        key_paragraphs = [p for p in key.paragraphs if re.match(r"^\d+\. ", p.text)]
        expected_numbers = [number for section in fmt["sections"] for number in range(1, section["total_questions"] + 1)]
        self.assertEqual([int(p.text.split(".", 1)[0]) for p in question_paragraphs], expected_numbers)
        self.assertEqual([int(p.text.split(".", 1)[0]) for p in key_paragraphs], expected_numbers)
        self.assertEqual(len(question_paragraphs), question_count)
        self.assertTrue(all("\t[" in p.text for p in question_paragraphs))
        self.assertEqual(
            sum(int(re.search(r"\[(\d+)\]$", p.text).group(1)) for p in question_paragraphs),
            total_marks,
        )
        for paragraph in question_paragraphs:
            stops = list(paragraph.paragraph_format.tab_stops)
            self.assertTrue(any(stop.alignment == WD_TAB_ALIGNMENT.RIGHT for stop in stops))
        self.assertNotIn("Answer not entered", "\n".join(p.text for p in key.paragraphs))
        self.assertNotIn("________________", "\n".join(p.text for p in paragraphs))

        index = 0
        all_text = "\n".join(p.text for p in paragraphs)
        self.assertIn(f"Maximum Marks: {total_marks}", all_text)
        self.assertIn(fmt["school_name"], all_text)
        for section in fmt["sections"]:
            self.assertIn(section["title"], all_text)
            for number in range(1, section["total_questions"] + 1):
                question = restored.get_question(section["id"], number)
                paragraph = question_paragraphs[index]
                self.assertIn(question["text"], paragraph.text)
                self.assertTrue(paragraph.text.endswith(f"[{section['marks_per_question']}]"))
                if section["question_type"] == "MCQ":
                    answer_label = "abcd"[question["options"].index(question["answer"])]
                    self.assertEqual(key_paragraphs[index].text, f"{number}. ({answer_label}) {question['answer']}")
                else:
                    self.assertEqual(key_paragraphs[index].text, f"{number}. {question['answer']}")
                if section["question_type"] == "MCQ":
                    paper_index = paragraphs.index(paragraph)
                    self.assertEqual(
                        [p.text for p in paragraphs[paper_index + 1:paper_index + 5]],
                        [f"({label}) {option}" for label, option in zip("abcd", question["options"])],
                    )
                if question["image_path"]:
                    paper_index = paragraphs.index(paragraph)
                    if question["image_position"] == "alongside":
                        self.assertTrue(paragraph._p.xpath(".//w:drawing"))
                    else:
                        self.assertTrue(paragraphs[paper_index + 1]._p.xpath(".//w:drawing"))
                index += 1

        expected_parts = sum(
            len(restored.get_question(section["id"], number)["sub_parts"])
            for section in fmt["sections"]
            for number in range(1, section["total_questions"] + 1)
        )
        actual_parts = [p for p in paragraphs if re.match(r"^\((i|ii)\) ", p.text)]
        self.assertEqual(len(actual_parts), expected_parts)

        with zipfile.ZipFile(paper_path) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(media), {25: 1, 50: 1, 100: 2}[total_marks])
            self.assertFalse(any(name.lower().endswith(".webp") for name in media))

    def test_incomplete_and_wrong_marks_warning(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fmt, state = build_scenario(25, root / "images")
            state.delete_question("A", 1)
            window = MainWindow({"paths": {
                "formats_dir": str(root / "formats"),
                "drafts_dir": str(root / "drafts"),
                "output_dir": str(root),
            }, "editor": {"auto_save_minutes": 2}})
            window._show_editor(fmt, state)
            with patch("ui.main_window.QMessageBox.question") as ask:
                from PyQt6.QtWidgets import QMessageBox

                ask.return_value = QMessageBox.StandardButton.Cancel
                self.assertFalse(window._confirm_incomplete())
                message = ask.call_args.args[2]
                self.assertIn("1 question(s) are empty", message)
                self.assertIn("24 marks", message)
            window.close()

    def test_missing_image_is_not_silently_omitted(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fmt, state = build_scenario(25, root / "images")
            Path(state.get_question("B", 2)["image_path"]).unlink()
            target = root / "broken.docx"
            with self.assertRaises(ImageError):
                DocumentEngine(fmt, state).save(str(target))
            self.assertFalse(target.exists())
            pdf_target = root / "broken.pdf"
            with self.assertRaises(ImageError):
                PdfEngine(fmt, state).save(str(pdf_target))
            self.assertFalse(pdf_target.exists())
