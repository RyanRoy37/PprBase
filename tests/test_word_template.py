"""A Word sample becomes a reusable, empty paper format."""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PIL import Image
from PyQt6.QtWidgets import QApplication

from core.document_engine import DocumentEngine
from core.draft_store import DraftStore
from core.format_loader import load_format
from core.state_manager import PaperState
from core.word_template import inspect_word_template, install_word_template
from ui.format_selector import FormatSelector
from ui.format_wizard import FormatWizard, update_format

ROOT = Path(__file__).resolve().parent / "fixtures" / "chemistry"
SOURCES = {
    "VI": ROOT / "Chemistry Ut1 VI Blueprint.docx",
    "VII": ROOT / "Chemistry Ut1 VII qs paper 2024.docx",
}


class WordTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def test_both_samples_become_reusable_formats(self) -> None:
        for klass, old_question in (("VI", "Risa prepared lemonade"), ("VII", "Vishal has a mixture")):
            with self.subTest(klass=klass), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                source = root / "uploaded.docx"
                shutil.copy2(SOURCES[klass], source)
                original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
                proposal = inspect_word_template(source)
                self.assertEqual(proposal["template_body_start"], 4)
                self.assertEqual(proposal["total_marks"], 25)
                self.assertEqual([s["total_questions"] for s in proposal["sections"]], [6, 3, 3, 3, 3, 3])
                self.assertEqual([s["marks_per_question"] for s in proposal["sections"]], [0.5, 1, 1, 1, 1, 5])
                self.assertEqual(proposal["sections"][-1]["attempt_count"], 2)

                format_path = install_word_template(proposal, source, root / "formats")
                fmt = load_format(format_path)
                self.assertNotEqual(Path(fmt["docx_template"]), source)
                self.assertTrue(Path(fmt["docx_template"]).is_file())
                source.unlink()  # The reusable template must outlive the upload.

                state = PaperState(fmt["format_id"])
                state.add_question("A", 1, {"text": "Which element is a metal?", "options": ["Iron", "Oxygen", "Carbon", "Sulphur"], "marks": 0.5, "type": "MCQ"})
                paper_path = root / "new_paper.docx"
                DocumentEngine(fmt, state).save(str(paper_path))
                paper = Document(paper_path)
                text = "\n".join(p.text for p in paper.paragraphs)
                self.assertIn("Which element is a metal?", text)
                self.assertNotIn(old_question, text)
                self.assertIn("St Vincent Pallotti School", text)
                self.assertIn("Unit Test-1", text)
                self.assertAlmostEqual(paper.sections[0].left_margin.cm, 1.27, places=2)
                self.assertAlmostEqual(paper.sections[0].right_margin.cm, 1.27, places=2)
                self.assertEqual(paper.styles["Normal"].font.size.pt, 12)
                first_new_heading = next(p for p in paper.paragraphs if p.text.startswith("Section- 1 —"))
                self.assertEqual(first_new_heading.runs[0].font.size.pt, 14)
                self.assertEqual(len(paper.inline_shapes), 0)
                self.assertEqual(original_hash, hashlib.sha256(SOURCES[klass].read_bytes()).hexdigest())

                store = DraftStore(str(root / "drafts"))
                draft_id = store.save(fmt, state)
                Path(fmt["docx_template"]).unlink()
                restored_fmt, restored_state = store.load(draft_id)
                self.assertTrue(Path(restored_fmt["docx_template"]).is_file())
                DocumentEngine(restored_fmt, restored_state).save(str(root / "resumed_paper.docx"))
                self.assertIn("Which element is a metal?", "\n".join(p.text for p in Document(root / "resumed_paper.docx").paragraphs))

    def test_other_word_layout_is_inferred_from_its_own_headings(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "History Quiz.docx"
            doc = Document()
            doc.add_paragraph("Sample Academy")
            doc.add_paragraph("History Quiz")
            doc.add_paragraph("Subject-History")
            doc.add_paragraph("Class-IX   Marks:14")
            doc.add_paragraph("Section A")
            doc.add_paragraph("Short answers (4x2=8)")
            doc.add_paragraph("Section B")
            doc.add_paragraph("Explain briefly (2x3=6)")
            doc.save(path)
            fmt = inspect_word_template(path)
            self.assertEqual(fmt["subject"], "History")
            self.assertEqual(fmt["class"], "IX")
            self.assertEqual(fmt["total_marks"], 14)
            self.assertEqual([(s["total_questions"], s["marks_per_question"]) for s in fmt["sections"]], [(4, 2), (2, 3)])

    def test_header_image_stays_and_old_question_image_is_removed(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            logo = root / "logo.png"
            old_diagram = root / "old_diagram.png"
            Image.new("RGB", (80, 40), "red").save(logo)
            Image.new("RGB", (80, 40), "blue").save(old_diagram)
            source = root / "sample.docx"
            doc = Document()
            doc.add_paragraph("Example School")
            doc.add_paragraph().add_run().add_picture(str(logo))
            doc.add_paragraph("Subject-Science")
            doc.add_paragraph("Class-5 Marks:2")
            doc.add_paragraph("Section A")
            doc.add_paragraph("Choose one (2x1=2)")
            doc.add_paragraph("Old question about the blue diagram")
            doc.add_paragraph().add_run().add_picture(str(old_diagram))
            doc.save(source)
            proposal = inspect_word_template(source)
            path = install_word_template(proposal, source, root / "formats")
            fmt = load_format(path)
            state = PaperState(fmt["format_id"])
            state.add_question("A", 1, {"text": "New question about plants?", "marks": 1, "type": "MCQ"})
            output = root / "new.docx"
            DocumentEngine(fmt, state).save(str(output))
            generated = Document(output)
            self.assertEqual(len(generated.inline_shapes), 1)
            text = "\n".join(p.text for p in generated.paragraphs)
            self.assertIn("New question about plants?", text)
            self.assertNotIn("Old question about the blue diagram", text)

    def test_plain_numbered_questions_can_be_reviewed(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "plain.docx"
            doc = Document()
            doc.add_paragraph("Local School")
            doc.add_paragraph("Term Test")
            doc.add_paragraph("Subject-Geography")
            doc.add_paragraph("Class-6 Marks:8")
            doc.add_paragraph("Question 1: Name a river.")
            doc.add_paragraph("Question 2: Name a mountain.")
            doc.save(path)
            fmt = inspect_word_template(path)
            self.assertEqual(fmt["template_body_start"], 4)
            self.assertEqual(len(fmt["sections"]), 1)
            self.assertEqual(fmt["sections"][0]["total_questions"], 2)

    def test_edits_to_imported_header_and_instructions_are_exported(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            proposal = inspect_word_template(SOURCES["VI"])
            proposal.update({
                "school_name": "Another School",
                "paper_title": "Practice Test",
                "subject": "General Science",
                "class": "VIII",
                "duration": "40 minutes",
                "instructions": ["Write your name on every page."],
            })
            fmt = load_format(install_word_template(proposal, SOURCES["VI"], root / "formats"))
            output = root / "changed_header.docx"
            DocumentEngine(fmt, PaperState(fmt["format_id"])).save(str(output))
            exported = "\n".join(p.text for p in Document(output).paragraphs)
            for phrase in ("Another School", "Practice Test", "Subject-General Science", "Class-VIII", "Time: 40 minutes", "Write your name on every page."):
                self.assertIn(phrase, exported)
            self.assertNotIn("St Vincent Pallotti School", exported)

    def test_import_button_creates_a_format_card(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            selector = FormatSelector(str(Path(temporary) / "formats"))
            chosen = []
            selector.format_chosen.connect(lambda fmt, path: chosen.append((fmt, path)))
            with patch("ui.format_selector.QFileDialog.getOpenFileName", return_value=(str(SOURCES["VI"]), "")), patch(
                "ui.format_selector.FormatWizard.exec", return_value=1
            ):
                selector._import_word_template()
            self.assertEqual(len(chosen), 1)
            self.assertTrue(Path(chosen[0][1]).is_file())
            self.assertTrue(Path(chosen[0][0]["docx_template"]).is_file())
            selector.close()

    def test_saved_template_is_reusable_and_editable_after_reopening(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = install_word_template(inspect_word_template(SOURCES["VI"]), SOURCES["VI"], root / "formats")
            first = load_format(path)
            for question, name in (("Why does ice float?", "first.docx"), ("What is an atom?", "second.docx")):
                reopened = load_format(path)
                state = PaperState(reopened["format_id"])
                self.assertEqual(state.count_filled("A"), 0)
                state.add_question("B", 1, {"text": question, "marks": 1})
                DocumentEngine(reopened, state).save(str(root / name))
            first_text = "\n".join(p.text for p in Document(root / "first.docx").paragraphs)
            second_text = "\n".join(p.text for p in Document(root / "second.docx").paragraphs)
            self.assertIn("Why does ice float?", first_text)
            self.assertNotIn("Why does ice float?", second_text)
            self.assertIn("What is an atom?", second_text)

            edited = dict(first)
            edited["display_name"] = "My reusable chemistry paper"
            edited["sections"] = [dict(section) for section in first["sections"]]
            edited["sections"][0]["title"] = "My multiple choice section"
            update_format(edited, path)
            reopened = load_format(path)
            self.assertEqual(reopened["display_name"], "My reusable chemistry paper")
            self.assertEqual(reopened["sections"][0]["title"], "My multiple choice section")
            DocumentEngine(reopened, PaperState(reopened["format_id"])).save(str(root / "after_edit.docx"))
            self.assertIn("My multiple choice section", "\n".join(p.text for p in Document(root / "after_edit.docx").paragraphs))

    def test_replacement_word_file_updates_future_papers_and_preserves_draft_source(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = install_word_template(inspect_word_template(SOURCES["VI"]), SOURCES["VI"], root / "formats")
            original = load_format(path)
            draft_id = DraftStore(str(root / "drafts")).save(original, PaperState(original["format_id"]))

            replacement = root / "new_sample.docx"
            sample = Document()
            for line in (
                "New Academy", "Semester Exam", "Subject-Physics", "Class-VIII Marks:10",
                "Section A", "Short answers (5x2=10)", "Old question to remove",
            ):
                sample.add_paragraph(line)
            sample.save(replacement)
            selector = FormatSelector(str(root / "formats"))
            def choose_replacement(dialog):
                dialog._replace_word_file()
                return 1
            with patch("ui.format_wizard.QFileDialog.getOpenFileName", return_value=(str(replacement), "")), patch.object(
                FormatWizard, "exec", new=choose_replacement
            ):
                selector._edit_format(path)
            selector.close()
            replacement.unlink()
            changed = load_format(path)
            self.assertEqual(changed["format_id"], original["format_id"])
            self.assertNotEqual(changed["docx_template"], original["docx_template"])
            self.assertEqual(changed["total_marks"], 10)
            self.assertEqual(changed["sections"][0]["total_questions"], 5)
            self.assertTrue(Path(changed["docx_template"]).is_file())
            state = PaperState(changed["format_id"])
            state.add_question("A", 1, {"text": "Explain reflection of light.", "marks": 2})
            output = root / "revised.docx"
            DocumentEngine(changed, state).save(str(output))
            text = "\n".join(p.text for p in Document(output).paragraphs)
            self.assertIn("New Academy", text)
            self.assertIn("Explain reflection of light.", text)
            self.assertNotIn("Old question to remove", text)

            old_fmt, old_state = DraftStore(str(root / "drafts")).load(draft_id)
            old_output = root / "old_draft.docx"
            DocumentEngine(old_fmt, old_state).save(str(old_output))
            old_text = "\n".join(p.text for p in Document(old_output).paragraphs)
            self.assertIn("St Vincent Pallotti School", old_text)
            self.assertNotIn("New Academy", old_text)


if __name__ == "__main__":
    unittest.main()
