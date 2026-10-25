"""Regression checks for editing, recovery and Word output."""

import os
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PIL import Image
from PyQt6.QtWidgets import QApplication

from core.document_engine import DocumentEngine
from core.draft_store import DraftStore
from core.format_loader import FormatError, load_format, validate_format
from core.state_manager import PaperState
from ui.paper_editor import PaperEditor
from ui.format_wizard import FormatWizard, save_new_format, update_format
from ui.main_window import MainWindow
import main


ROOT = Path(__file__).resolve().parents[1]
FORMAT = load_format(str(ROOT / "formats" / "class7_chemistry.json"))


class WorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_navigation_draft_and_export(self):
        state = PaperState(FORMAT["format_id"])
        editor = PaperEditor(FORMAT, state)
        editor.show()
        editor.form.text_edit.setPlainText("What is an atom?")
        editor.form.option_edits[0].setText("A particle")
        editor.form.option_edits[1].setText("A planet")
        editor.form.answer_combo.setCurrentIndex(1)
        editor.q_spin.setValue(2)
        self.assertEqual(state.get_question("A", 1)["text"], "What is an atom?")
        editor.form.text_edit.setPlainText("Choose the element.")
        editor.save_current()

        with tempfile.TemporaryDirectory() as temp:
            store = DraftStore(str(Path(temp) / "drafts"))
            draft_id = store.save(FORMAT, state)
            restored_format, restored = store.load(draft_id)
            self.assertEqual(restored.get_question("A", 1)["answer"], "A particle")
            self.assertEqual(restored.get_question("A", 2)["text"], "Choose the element.")
            self.assertEqual(store.list_drafts()[0]["draft_id"], draft_id)

            paper_path = Path(temp) / "paper.docx"
            engine = DocumentEngine(restored_format, restored)
            engine.save(str(paper_path))
            paper = Document(paper_path)
            content = "\n".join(p.text for p in paper.paragraphs)
            self.assertIn("What is an atom?", content)
            self.assertIn("A particle", content)
            self.assertIn("Maximum Marks: 70", content)
            self.assertNotIn("Answer Key", content)

            key_path = Path(temp) / "answers.docx"
            engine.build_answer_key().save(key_path)
            key = Document(key_path)
            key_text = "\n".join(p.text for p in key.paragraphs)
            self.assertIn("(a) A particle", key_text)
        editor.close()

    def test_subpart_navigation_preserves_main_question(self):
        state = PaperState(FORMAT["format_id"])
        editor = PaperEditor(FORMAT, state)
        editor.show()
        editor.tab_bar.setCurrentIndex(1)
        editor.form.text_edit.setPlainText("Explain the reaction.")
        editor.subpart_combo.setCurrentIndex(1)
        editor.form.text_edit.setPlainText("Name the reactants.")
        editor.q_spin.setValue(2)
        question = state.get_question("B", 1)
        self.assertEqual(question["text"], "Explain the reaction.")
        self.assertEqual(question["sub_parts"]["i"]["text"], "Name the reactants.")
        self.assertEqual(state.allocated_marks(FORMAT), 2)
        editor.close()

    def test_format_creation_and_resume_from_home(self):
        with tempfile.TemporaryDirectory() as temp:
            wizard = FormatWizard()
            wizard.subject.setText("Biology")
            wizard.klass.setText("6")
            wizard.sections.append({
                "id": "A", "title": "Section A", "description": "Answer all",
                "question_type": "short_answer", "marks_per_question": 2,
                "total_questions": 3, "sub_parts": None,
            })
            fmt_path = save_new_format(wizard.format_data(), str(Path(temp) / "formats"))
            fmt = load_format(fmt_path)
            self.assertEqual(fmt["total_marks"], 6)
            fmt["school_name"] = "Sample School"
            update_format(fmt, fmt_path)
            self.assertEqual(load_format(fmt_path)["school_name"], "Sample School")

            window = MainWindow({"paths": {
                "formats_dir": str(Path(temp) / "formats"),
                "drafts_dir": str(Path(temp) / "drafts"),
                "output_dir": str(Path(temp) / "output"),
            }, "editor": {"auto_save_minutes": 2}})
            self.assertEqual(window.windowTitle(), "PprBase")
            window.show()
            window._open_format(fmt, fmt_path)
            window.editor.form.text_edit.setPlainText("What is a cell?")
            window._go_home()
            drafts = window.drafts.list_drafts()
            self.assertEqual(len(drafts), 1)
            window._resume_draft(drafts[0]["draft_id"])
            self.assertEqual(window.editor.form.text_edit.toPlainText(), "What is a cell?")
            window.close()
            wizard.close()

    def test_draft_keeps_attached_image(self):
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "diagram.png"
            Image.new("RGB", (100, 100), "white").save(image_path)
            state = PaperState(FORMAT["format_id"])
            state.add_question("A", 1, {"text": "See diagram", "type": "MCQ", "marks": 1,
                                        "image_path": str(image_path)})
            store = DraftStore(str(Path(temp) / "drafts"))
            draft_id = store.save(FORMAT, state)
            image_path.unlink()
            fmt, restored = store.load(draft_id)
            self.assertTrue(Path(restored.get_question("A", 1)["image_path"]).is_file())
            output = Path(temp) / "paper.docx"
            DocumentEngine(fmt, restored).save(str(output))
            self.assertEqual(len(Document(output).inline_shapes), 1)

    def test_packaged_startup_uses_user_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp) / "bundle"
            (bundle / "config").mkdir(parents=True)
            (bundle / "formats").mkdir()
            (bundle / "config" / "settings.json").write_text(json.dumps({
                "paths": {"formats_dir": "formats", "drafts_dir": "drafts", "output_dir": "output"}
            }), encoding="utf-8")
            (bundle / "formats" / "sample.json").write_text(json.dumps(FORMAT), encoding="utf-8")
            app_data = str(Path(temp) / "AppData")
            with patch.object(main, "CONFIG_PATH", str(bundle / "config" / "settings.json")), \
                 patch.object(main, "BUNDLE_DIR", str(bundle)), \
                 patch.object(sys, "frozen", True, create=True), \
                 patch.dict(os.environ, {"APPDATA": app_data}):
                config = main.load_config()
            self.assertTrue(Path(config["paths"]["formats_dir"], "sample.json").is_file())
            self.assertTrue(Path(config["paths"]["drafts_dir"]).is_dir())
            self.assertTrue(config["paths"]["output_dir"].startswith(app_data))
            self.assertEqual(Path(config["paths"]["formats_dir"]).parent.name, "PprBase")

    def test_old_app_data_and_draft_images_migrate(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = root / "bundle"
            (bundle / "config").mkdir(parents=True)
            (bundle / "formats").mkdir()
            (bundle / "config" / "settings.json").write_text(json.dumps({
                "paths": {"formats_dir": "formats", "drafts_dir": "drafts", "output_dir": "output"}
            }), encoding="utf-8")
            (bundle / "formats" / "sample.json").write_text(json.dumps(FORMAT), encoding="utf-8")

            app_data = root / "AppData"
            legacy = app_data / "QuestionCraft"
            (legacy / "formats").mkdir(parents=True)
            (legacy / "formats" / "old.json").write_text(json.dumps(FORMAT), encoding="utf-8")
            source_image = root / "diagram.png"
            Image.new("RGB", (100, 100), "white").save(source_image)
            state = PaperState(FORMAT["format_id"])
            state.add_question("A", 1, {"text": "See the image", "marks": 1,
                                        "image_path": str(source_image)})
            draft_id = DraftStore(str(legacy / "drafts")).save(FORMAT, state)

            with patch.object(main, "CONFIG_PATH", str(bundle / "config" / "settings.json")), \
                 patch.object(main, "BUNDLE_DIR", str(bundle)), \
                 patch.object(sys, "frozen", True, create=True), \
                 patch.dict(os.environ, {"APPDATA": str(app_data)}):
                config = main.load_config()

            self.assertTrue((app_data / "PprBase" / "formats" / "old.json").is_file())
            self.assertTrue(legacy.is_dir())
            legacy.rename(app_data / "QuestionCraft_backup")
            _, restored = DraftStore(config["paths"]["drafts_dir"]).load(draft_id)
            self.assertTrue(Path(restored.get_question("A", 1)["image_path"]).is_file())

    def test_bad_format_reports_an_error(self):
        bad = dict(FORMAT)
        bad["sections"] = [{**FORMAT["sections"][0], "question_type": []}]
        with self.assertRaises(FormatError):
            validate_format(bad)


if __name__ == "__main__":
    unittest.main()
